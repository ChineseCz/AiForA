from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage, PageBreak, KeepTogether

OUT = Path('reports/gold_research_2026')
OUT.mkdir(parents=True, exist_ok=True)
IMG = OUT / 'figures'
IMG.mkdir(exist_ok=True)
FONT = 'C:/Windows/Fonts/simhei.ttf'

sources = [
    ('世界黄金协会：《2026 Central Bank Gold Reserves Survey》', 'https://www.gold.org/goldhub/research/central-bank-gold-reserves-survey-2026'),
    ('世界黄金协会：《Gold Demand Trends Q2 2026: Central Banks》', 'https://www.gold.org/goldhub/research/gold-demand-trends/gold-demand-trends-q2-2026/central-banks'),
    ('世界黄金协会：《Gold Demand Trends Q2 2026: Outlook》', 'https://www.gold.org/goldhub/research/gold-demand-trends/gold-demand-trends-q2-2026/outlook'),
    ('世界黄金协会：《Gold Mid-Year Outlook 2026》', 'https://www.gold.org/goldhub/research/gold-mid-year-outlook-2026'),
    ('Reuters：《Gold forecasts fall, but central bank buying expected to cushion retreat》', 'https://www.reuters.com/world/india/gold-forecasts-fall-central-bank-buying-expected-cushion-retreat-2026-07-28/'),
    ('Bloomberg：China Central Bank Adds Most Gold Since 2023 Even as Prices Jump', 'https://www.bloomberg.com/news/articles/2026-09-07/china-central-bank-adds-most-gold-since-2023-even-as-prices-jump'),
]

def fnt(size, bold=False):
    return ImageFont.truetype(FONT, size)

def chart_bar(path, title, labels, values, suffix='', colors_=None, note=''):
    W,H=1400,760; im=Image.new('RGB',(W,H),'white'); d=ImageDraw.Draw(im)
    d.text((60,35), title, font=fnt(36, True), fill='#172033')
    left, bottom, top, right = 130, 650, 145, 130
    maxv=max(values)*1.25
    barw=(W-left-right)//(len(values)*2)
    gap=barw
    cols=colors_ or ['#C99A2E']*len(values)
    for i,(lab,val) in enumerate(zip(labels,values)):
        x=left+gap//2+i*2*barw
        y=bottom-int((bottom-top)*val/maxv)
        d.rectangle((x,y,x+barw,bottom), fill=cols[i])
        d.text((x+barw//2-35,y-42), f'{val}{suffix}', font=fnt(28, True), fill='#172033')
        d.text((x+barw//2-100,bottom+20), lab, font=fnt(25), fill='#334155')
    d.line((left,top,left,bottom), fill='#9CA3AF', width=2); d.line((left,bottom,W-right,bottom), fill='#9CA3AF', width=2)
    if note: d.text((60,700),note,font=fnt(20),fill='#64748B')
    im.save(path)

def chart_scenario(path):
    W,H=1600,900; im=Image.new('RGB',(W,H),'white'); d=ImageDraw.Draw(im)
    d.text((60,35),'黄金未来走势：情景分析框架（非价格预测）',font=fnt(36,True),fill='#172033')
    rows=[('偏多','经济放缓、实际利率下降、美元走弱、风险升级','上行空间打开','#DDF5E5'),('基准','增长温和、实际利率高位、央行继续购金','高位震荡','#FFF4CC'),('偏空','增长强劲、实际利率上升、美元走强、ETF流出','估值消化','#FDE2E2')]
    y=150
    for name,cond,outcome,bg in rows:
        d.rounded_rectangle((60,y,W-60,y+185),20,fill=bg,outline='#CBD5E1',width=2)
        d.text((95,y+35),name,font=fnt(34,True),fill='#172033')
        d.text((350,y+35),cond,font=fnt(27),fill='#334155')
        d.text((350,y+105),outcome,font=fnt(30,True),fill='#172033')
        y+=220
    d.text((60,825),'资料：世界黄金协会2026年中期展望；本图用于帮助投资者建立概率思维。',font=fnt(20),fill='#64748B')
    im.save(path)

chart_bar(IMG/'central_bank_accumulation.png','全球央行年均购金量对比',["此前十年","过去四年"],[500,1000],'吨',['#94A3B8','#C99A2E'],'来源：世界黄金协会2026年央行储备调查。')
chart_bar(IMG/'central_bank_survey.png','2026年央行调查核心结果',["全球储备未来12个月增加","自身储备未来12个月增加","五年后黄金占比更高"],[89,45,84],'%',['#C99A2E','#D8B65A','#E7C979'],'来源：世界黄金协会2026年央行储备调查。')
chart_scenario(IMG/'scenario_matrix.png')

md = f'''# 黄金市场与国家队购金趋势研究报告

**报告日期：2026年9月7日**  
**研究对象：全球黄金、中国人民银行购金、普通投资者资产配置**  
**报告性质：公开资料研究，不构成个性化投资建议**

> **核心结论**：全球央行购金仍处于历史高位，黄金的长期储备与避险逻辑没有破坏；但黄金已经历大幅上涨和剧烈回撤，当前更适合把它作为组合中的防御性资产分批配置，而不是因为“国家队在买”而一次性满仓。

## 1. 投资摘要

### 1.1 国家队趋势仍偏多，但不是短线买入信号

世界黄金协会数据显示，过去四年全球央行平均每年净购金约 **1000吨**，而此前十年平均约 **500吨**。2026年调查中，89%的受访央行预计未来12个月全球央行黄金储备继续增加，45%预计自身储备增加，74%预计未来五年全球储备中的美元持有量下降。

中国人民银行方面，世界黄金协会披露：2026年第二季度增持约33吨，上半年增持约40吨，截至6月底官方黄金储备约2346吨。9月7日市场报道显示，中国央行8月继续加快购金，但单月数据应以官方披露和后续修订为准。

### 1.2 价格逻辑从“单一避险”变为多因素博弈

黄金未来走势取决于：美国实际利率、美元、ETF资金流、央行购金、地缘政治、人民币汇率和亚洲实物需求。央行购金提供结构性支撑，但短期价格经常由实际利率和资金仓位主导。

### 1.3 当前不是低风险追高位置

黄金在2026年初曾超过5500美元/盎司，6月底一度跌破4000美元/盎司，9月初回到4400美元/盎司附近。路透社7月底调查的29位分析师将2026年黄金平均预测中位数下调至约4509美元/盎司，说明市场仍偏多，但对短期上行空间更谨慎。

## 2. 数据概览

![全球央行年均购金量](figures/central_bank_accumulation.png)

![央行调查结果](figures/central_bank_survey.png)

### 2.1 关键数据表

| 指标 | 最新公开信息 | 研究含义 |
|---|---:|---|
| 全球央行过去四年年均购金 | 约1000吨 | 明显高于此前十年均值，储备多元化趋势增强 |
| 全球央行未来12个月储备预期 | 89%预计增加 | 官方部门的战略意愿仍强 |
| 受访央行预计自身增持 | 45% | 购金主体仍有扩散可能 |
| 受访央行预计五年后黄金占比提高 | 84% | 黄金在储备管理中的地位上升 |
| 全球央行2026年Q2净购金 | 约289吨 | Q2明显反弹，但H1仍受Q1卖出影响 |
| 中国央行2026年Q2增持 | 约33吨 | 中国购金仍然持续且具战略属性 |
| 中国央行2026年H1增持 | 约40吨 | 速度不是直线上升，但方向未改变 |

## 3. 为什么央行持续买黄金？

### 3.1 外汇储备多元化

黄金不依赖某个国家或机构兑付，不是任何发行人的负债。对于央行来说，黄金可以降低对单一货币、单一金融体系和单一托管体系的依赖。

### 3.2 对冲地缘政治与金融制裁风险

黄金在国内储存时，不同于海外金融资产，不容易受到单一司法管辖区的冻结影响。需要强调的是，这体现的是国家资产负债表的风险管理，不等同于普通投资者的短线交易信号。

### 3.3 对冲通胀和货币信用风险

在财政赤字扩大、债务持续累积、货币购买力存在不确定性时，黄金可作为长期储备资产。但黄金并不产生利息，持有它的收益来自价格变化和货币变化。

## 4. 黄金价格驱动框架

### 4.1 美国实际利率

实际利率上升，持有无息黄金的机会成本提高，通常对金价不利；实际利率下降，黄金相对吸引力增强。判断时应看实际利率，而不是只看名义利率。

### 4.2 美元指数

国际黄金以美元计价。美元强势通常压制黄金，美元弱势通常支持黄金，但危机期间美元和黄金可能同时受益于避险需求。

### 4.3 ETF资金流

央行购金更像结构性支撑，黄金ETF资金流更容易影响短期趋势。央行继续买而ETF流出，黄金可能高位震荡；央行继续买且ETF重新流入，趋势行情更容易延续。

### 4.4 中国人民币汇率

人民币投资者实际承担的是人民币金价风险。粗略关系为：

> 人民币金价 ≈ 国际美元金价 × 美元兑人民币汇率 + 国内溢价/折价

因此，国际金价不涨时，人民币贬值也可能推高国内金价；反之，人民币升值可能压低国内金价涨幅。

## 5. 未来走势：三种情景

![黄金情景分析](figures/scenario_matrix.png)

| 情景 | 主要条件 | 可能结果 | 投资含义 |
|---|---|---|---|
| 偏多 | 经济放缓、实际利率下降、美元走弱、地缘风险升级 | 黄金重新走强，甚至测试前高 | 适合已有配置者持有，但不宜追涨加杠杆 |
| 基准 | 经济温和、实际利率高位、央行继续购金但放缓 | 高位宽幅震荡 | 分批配置、目标比例再平衡 |
| 偏空 | 经济强劲、实际利率上升、美元走强、ETF流出 | 5%—15%估值消化并不意外 | 避免一次性重仓，保留现金 |

世界黄金协会2026年中期展望认为，在宏观预期不显著改变的情况下，黄金可能处于相对震荡状态；若经济恶化、降息预期增强或地缘冲突升级，黄金可能重新上行；若增长强劲、收益率上升、市场风险偏好恢复，则黄金可能回调。

## 6. 现在是否值得买？

### 6.1 如果目标是短线赚钱

不建议把黄金视为低风险短线工具。当前价格已经经历大幅上涨和回撤，短期波动受利率和资金仓位影响很大，不能因为央行购金就推断未来几个月一定上涨。

### 6.2 如果目标是十年以上保值和分散风险

可以考虑配置。黄金更适合承担“保险”和“组合分散”角色，而不是承担主要收益来源。

### 6.3 如果没有应急资金或存在高息负债

优先建立现金储备和偿还高息债务，不应先买黄金。

## 7. 普通投资者的配置方法

### 7.1 目标比例

对多数非专业投资者，可以把黄金作为总金融资产的 **5%—10%** 参考配置区间：

- 0%—3%：对组合影响很小；
- 5%—10%：具备一定对冲和分散效果；
- 10%—15%：已经包含较明显的黄金观点；
- 超过20%：更像方向性押注，不再是普通意义上的分散配置。

### 7.2 购买方式

**黄金ETF/黄金基金**：透明、易交易、无需保管，适合长期配置；关注管理费、跟踪误差和溢价折价。  
**银行积存金**：适合定投，但要比较买卖价差和实物提取成本。  
**投资金条/金币**：有实物属性，但需考虑保管、回收渠道和买卖价差。  
**黄金首饰**：主要是消费，不宜作为主要投资工具，因为工费和回收折价较高。

### 7.3 建议分批，而非一次性买入

如果计划配置总资产的8%，可以分3—6次投入，或采用每月固定金额定投。目标不是精准预测最低点，而是降低买在短期情绪高点的风险。

## 8. 行外人投资决策清单

在买入前依次回答：

1. 我买黄金是为了保值、避险，还是为了短期赚钱？
2. 如果黄金短期下跌15%，我是否会被迫卖出？
3. 我的应急现金和保险是否已经准备好？
4. 黄金占我的全部金融资产比例是多少？
5. 我选择的产品买卖价差、管理费和流动性如何？
6. 我是否因为害怕错过上涨而冲动买入？
7. 我准备在什么情况下减仓或再平衡？

## 9. 每月跟踪的六项指标

1. 美国10年期实际利率；
2. 美元指数；
3. 全球黄金ETF资金流和持仓；
4. 全球央行季度购金量；
5. 中国人民银行黄金储备；
6. 美元兑人民币汇率与国内黄金溢价。

## 10. 风险提示

- 黄金价格可能出现较大回撤；
- 央行可能因流动性需要阶段性卖出；
- 实际利率和美元走强可能压制金价；
- ETF可能持续流出；
- 实物黄金存在保管、鉴定和回收风险；
- 国内金价还受到人民币汇率和本地溢价影响；
- 不应借贷或使用高杠杆购买黄金；
- 历史上涨不代表未来收益。

## 11. 研究结论

黄金的长期结构性逻辑仍然成立：全球央行购金、储备多元化、地缘政治不确定性和财政债务压力，都为黄金提供了长期支撑。但黄金已经不是低估值、低波动的资产，未来收益率可能低于过去两年的表现，波动率可能继续偏高。

**最终判断：**

> 对普通投资者，黄金值得作为5%—10%左右的长期防御性配置进行分批布局；不值得因为“国家队在买”而一次性满仓，更不适合借钱、加杠杆或把黄金当成短期暴涨彩票。

## 参考资料

'''
for title,url in sources:
    md += f'- [{title}]({url})\n'
md += '\n*本报告数据截至2026年9月7日；部分央行数据可能存在披露滞后或后续修订。*\n'
(OUT/'黄金市场与国家队购金趋势研究报告.md').write_text(md, encoding='utf-8')

# DOCX

def set_cell_shading(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)

doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.78); sec.right_margin=Inches(.78)
styles=doc.styles; styles['Normal'].font.name='SimSun'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体'); styles['Normal'].font.size=Pt(10)
for name,size,color in [('Title',24,'1F2937'),('Heading 1',17,'9A6B16'),('Heading 2',13,'B07D1E')]:
    styles[name].font.name='SimHei'; styles[name]._element.rPr.rFonts.set(qn('w:eastAsia'),'黑体'); styles[name].font.size=Pt(size); styles[name].font.color.rgb=RGBColor.from_string(color); styles[name].font.bold=True
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('黄金市场与国家队购金趋势研究报告'); r.bold=True; r.font.size=Pt(24); r.font.name='SimHei'
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('全球央行购金、中国人民银行趋势与普通投资者配置框架').italic=True
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('报告日期：2026年9月7日  |  公开资料研究  |  非个性化投资建议')
doc.add_picture(str(IMG/'central_bank_accumulation.png'), width=Inches(6.4)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.add_heading('执行摘要', level=1)
for t in ['全球央行购金仍处于历史高位，但央行的长期储备行为不是普通投资者的短线买入信号。','中国人民银行购金方向仍偏战略性，节奏会波动，不能用单月数据预测金价。','黄金适合做组合中的防御性和分散性资产，不适合因为“国家队在买”而一次性满仓。','对多数普通投资者，可将黄金作为总金融资产5%—10%的参考配置，并分批买入。']:
    doc.add_paragraph(t, style='List Bullet')
doc.add_picture(str(IMG/'central_bank_survey.png'), width=Inches(6.4)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
# Use markdown headings/paragraphs in a compact DOCX conversion
sections=[]
for line in md.splitlines():
    if not line.strip() or line.startswith('![') or line.startswith('|') or line.startswith('- [') or line.startswith('>') or line.startswith('**核心结论**'):
        continue
    if line.startswith('# '): continue
    if line.startswith('## '): doc.add_heading(line[3:],1)
    elif line.startswith('### '): doc.add_heading(line[4:],2)
    elif line.startswith('- '): doc.add_paragraph(line[2:], style='List Bullet')
    elif line.startswith('1. ') or line.startswith('2. ') or line.startswith('3. ') or line.startswith('4. ') or line.startswith('5. ') or line.startswith('6. ') or line.startswith('7. '): doc.add_paragraph(line, style='List Number')
    elif line.startswith('> '): doc.add_paragraph(line[2:])
    elif line.strip().startswith('**') and line.strip().endswith('**'): doc.add_paragraph(line.strip('*'))
    elif not line.startswith('报告日期') and not line.startswith('说明：'):
        doc.add_paragraph(line)
doc.add_picture(str(IMG/'scenario_matrix.png'), width=Inches(6.4)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.add_heading('参考资料',1)
for title,url in sources: doc.add_paragraph(f'{title}：{url}', style='List Bullet')
doc.save(OUT/'黄金市场与国家队购金趋势研究报告.docx')

# PDF
pdfmetrics.registerFont(TTFont('SimHei', FONT))
pdf=SimpleDocTemplate(str(OUT/'黄金市场与国家队购金趋势研究报告.pdf'), pagesize=A4, rightMargin=16*mm,leftMargin=16*mm,topMargin=15*mm,bottomMargin=15*mm)
ss=getSampleStyleSheet(); normal=ParagraphStyle('CN',parent=ss['BodyText'],fontName='SimHei',fontSize=9.2,leading=15,spaceAfter=5); h1=ParagraphStyle('H1CN',parent=ss['Heading1'],fontName='SimHei',fontSize=16,textColor=colors.HexColor('#9A6B16'),spaceBefore=12,spaceAfter=7); h2=ParagraphStyle('H2CN',parent=ss['Heading2'],fontName='SimHei',fontSize=12,textColor=colors.HexColor('#B07D1E'),spaceBefore=8,spaceAfter=5); title=ParagraphStyle('TCN',parent=ss['Title'],fontName='SimHei',fontSize=22,leading=30,alignment=TA_CENTER,textColor=colors.HexColor('#1F2937')); small=ParagraphStyle('SmallCN',parent=normal,fontSize=7.5,textColor=colors.HexColor('#64748B'))
story=[Paragraph('黄金市场与国家队购金趋势研究报告',title),Spacer(1,4),Paragraph('全球央行购金、中国人民银行趋势与普通投资者配置框架',ParagraphStyle('sub',parent=normal,alignment=TA_CENTER,fontSize=11)),Paragraph('报告日期：2026年9月7日 | 公开资料研究 | 非个性化投资建议',ParagraphStyle('date',parent=small,alignment=TA_CENTER)),Spacer(1,10),RLImage(str(IMG/'central_bank_accumulation.png'),width=170*mm,height=92*mm),Spacer(1,8)]
for line in md.splitlines():
    if not line.strip() or line.startswith('![') or line.startswith('|') or line.startswith('- [') or line.startswith('# ') or line.startswith('报告日期') or line.startswith('说明：'):
        continue
    if line.startswith('## '): story.append(Paragraph(line[3:],h1))
    elif line.startswith('### '): story.append(Paragraph(line[4:],h2))
    elif line.startswith('- '): story.append(Paragraph('• '+line[2:],normal))
    elif line.startswith('> '): story.append(Paragraph('<b>'+line[2:]+'</b>',normal))
    elif line.startswith('**') and line.endswith('**'): story.append(Paragraph('<b>'+line.strip('*')+'</b>',normal))
    elif line.startswith('|'): continue
    elif line.startswith('1. ') or line.startswith('2. ') or line.startswith('3. ') or line.startswith('4. ') or line.startswith('5. ') or line.startswith('6. ') or line.startswith('7. '): story.append(Paragraph(line,normal))
    else: story.append(Paragraph(line.replace('&','&amp;'),normal))
story.extend([Spacer(1,6),RLImage(str(IMG/'central_bank_survey.png'),width=170*mm,height=92*mm),Spacer(1,6),RLImage(str(IMG/'scenario_matrix.png'),width=170*mm,height=96*mm)])
pdf.build(story)
print('created', OUT)
for p in sorted(OUT.glob('*')): print(p, p.stat().st_size)
